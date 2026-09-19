import collections
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def apply_text_style(run, font_name='Segoe UI', font_size=18, bold=False, color=RGBColor(255, 255, 255)):
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color

def add_title_slide(prs):
    slide_layout = prs.slide_layouts[0] # Title Slide
    slide = prs.slides.add_slide(slide_layout)
    
    # Set background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(11, 20, 26) # Dark background

    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "DynamicDetailing AI"
    apply_text_style(title.text_frame.paragraphs[0].runs[0], font_size=44, bold=True, color=RGBColor(0, 168, 132))

    subtitle.text = "Unified Voice & Chat CRM\nPowered by OmniDimension & n8n"
    for p in subtitle.text_frame.paragraphs:
        if p.runs:
            apply_text_style(p.runs[0], font_size=24, color=RGBColor(200, 200, 200))

def add_content_slide(prs, title_text, content_bullets):
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(11, 20, 26)

    title = slide.shapes.title
    content = slide.placeholders[1]

    title.text = title_text
    if title.text_frame.paragraphs[0].runs:
        apply_text_style(title.text_frame.paragraphs[0].runs[0], font_size=36, bold=True, color=RGBColor(0, 168, 132))

    content.text = ""
    for bullet in content_bullets:
        p = content.text_frame.add_paragraph()
        p.text = bullet
        apply_text_style(p.runs[0], font_size=20, color=RGBColor(233, 237, 239))
        p.space_before = Pt(14)
        p.space_after = Pt(14)
    
    # Remove the first empty paragraph
    if len(content.text_frame.paragraphs) > len(content_bullets):
        p = content.text_frame.paragraphs[0]
        p._element.getparent().remove(p._element)

def add_image_slide(prs, title_text, image_path, description):
    slide_layout = prs.slide_layouts[5] # Title Only
    slide = prs.slides.add_slide(slide_layout)
    
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(11, 20, 26)

    title = slide.shapes.title
    title.text = title_text
    if title.text_frame.paragraphs[0].runs:
        apply_text_style(title.text_frame.paragraphs[0].runs[0], font_size=36, bold=True, color=RGBColor(0, 168, 132))

    # Add Image
    try:
        left = Inches(1)
        top = Inches(2)
        height = Inches(4.5)
        slide.shapes.add_picture(image_path, left, top, height=height)
    except Exception as e:
        print(f"Could not load image: {e}")

    # Add Text box
    txBox = slide.shapes.add_textbox(Inches(5), Inches(2.5), Inches(4), Inches(3))
    tf = txBox.text_frame
    p = tf.add_paragraph()
    p.text = description
    apply_text_style(p.runs[0], font_size=24, color=RGBColor(233, 237, 239))

def main():
    prs = Presentation()
    
    # Set Slide Size (Widescreen 16:9)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(prs)

    add_content_slide(prs, "System Architecture", [
        "📞 Voice & Text: OmniDimension handles calls & WhatsApp natively.",
        "⚙️ Orchestration: n8n workflow background routing.",
        "📊 Backend: Google Sheets CRM Database.",
        "💻 Frontend: Next.js Dashboard (Phase 2)."
    ])

    img_path = r"C:\Users\SHAIK ATIF\.gemini\antigravity-ide\brain\5a51b896-3984-4a49-996d-955b502c0625\voice_call_active_mockup_1789231953475.jpg"
    add_image_slide(prs, "1. The Voice Engine: Code Switching", img_path, 
        "Hyper-realistic, code-switching conversational AI.\n\n"
        "The 'Human' Touch:\n"
        "- Uses Vyavaharikam (Spoken Telugu), not robotic text.\n"
        "- Code-switches to English for numbers and tech terms.\n"
        "- Latencies masked with conversational fillers."
    )

    add_content_slide(prs, "2. Text Engine: WhatsApp Menu", [
        "Interactive List Message bypassing raw text blocks.",
        "Instead of typing services, the AI sends a structured WhatsApp menu.",
        "Highlight: Paint Protection Film (PPF)",
        "- Provides a premium, app-like experience inside WhatsApp."
    ])

    add_content_slide(prs, "3. Text Engine: PPF Booking Handoff", [
        "Utility Message with Quick Reply Action Buttons.",
        "When a user books PPF, the AI confirms: 'Our manager will call you back.'",
        "Includes interactive buttons: [Cancel] [Call Me Now].",
        "Seamlessly routes high-ticket leads without friction."
    ])

    add_content_slide(prs, "4. CRM Orchestration (n8n)", [
        "Automated background switch logic via n8n Webhooks.",
        "HANDOFFS: Logged to Sheets + Meta API alerts manager via WhatsApp immediately.",
        "COMPLAINTS: Logged to Sheets + Flagged on Dashboard without buzzing manager directly."
    ])

    add_content_slide(prs, "Next Steps: Phase 2", [
        "Build Next.js React Dashboard",
        "Integrate Google Auth & Sheets API",
        "Migrate to Meta Cloud Business API",
        "Live Production Testing"
    ])

    prs.save("DynamicDetailing_AI_CRM.pptx")
    print("Presentation saved as DynamicDetailing_AI_CRM.pptx")

if __name__ == '__main__':
    main()
